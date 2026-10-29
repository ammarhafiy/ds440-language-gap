"""
Step 7: blind hand check of the automatic scores (Cohen's kappa).

Run after 06_score.py, once for each chatbot:
    # 1. Make the hand-check sheet: 200 replies, WITHOUT the automatic scores
    python scripts/07_hand_check.py make --scored data/scored/gemini_gemini-3.1-flash-lite_scored.csv

    # 2. Fill in the yellow "Your score" column yourself (rules on the "How to use" sheet
    #    and in notes/scoring_rubric.md), save the file in data/hand_check/

    # 3. Compare your scores with the automatic scores
    python scripts/07_hand_check.py kappa --scored data/scored/gemini_gemini-3.1-flash-lite_scored.csv

Reads:
    data/scored/<name>_scored.csv              scored replies from script 06
    data/questions/questions_final.csv         question text, facts and accepted answers
Writes:
    make:   data/hand_check/<name>_hand_check.xlsx
    kappa:  results/hand_check_<name>.csv              every sampled reply with both scores
            results/hand_check_summary_<name>.csv      agreement overall and per group, Cohen's kappa overall

Sample (notes/scoring_rubric.md, "Hand check"): 20 random replies from each of the 10 groups
(source x league x language x condition) = 200 replies, random seed 440. Failed requests are not
sampled. Within each group the replies are sorted by prompt_id before sampling, so the same seed
always gives the same sample (and two chatbots with no failed requests get the same 200 prompts).
The 200 rows are then shuffled so the groups are mixed on the sheet.

The sheet is blind: it shows the question, the fact given in the prompt (if any), the answer key,
the accepted answers and the reply, but not the automatic score.

Cohen's kappa = (p_o - p_e) / (1 - p_e), where p_o is the share of replies where my score and the
automatic score agree, and p_e is the agreement expected by chance from how often each scorer
says "correct". Kappa is not defined when p_e = 1 (both scorers gave the same score to every reply).
"""
import argparse
import csv
import random
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("This script needs openpyxl: pip install -r requirements.txt")

ROOT = Path(__file__).resolve().parent.parent
QUESTIONS_PATH = ROOT / "data" / "questions" / "questions_final.csv"
HAND_CHECK_DIR = ROOT / "data" / "hand_check"
RESULTS_DIR = ROOT / "results"

GROUP_COLUMNS = ["source", "league", "language", "condition"]
PER_GROUP = 20
SEED = 440
SHEET = "Hand check"
SCORE_HEADER = "Your score (correct / incorrect)"
ID_HEADER = "Prompt ID (do not edit)"

LEAGUE_NAMES = {"nba": "NBA", "msl": "Malaysia Super League", "mixed": "MKQA"}
LANGUAGE_NAMES = {"en": "English", "ms": "Malay"}
CONDITION_NAMES = {"no_facts": "no fact", "with_facts": "fact given"}

# Same look as notes/answer_key_spot_check.xlsx
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
ID_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------- helpers ----------

def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_csv(path, columns, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def name_of(scored_path):
    """data/scored/gemini_x_scored.csv -> gemini_x"""
    stem = Path(scored_path).stem
    return stem[: -len("_scored")] if stem.endswith("_scored") else stem


def group_label(row):
    if row["source"] == "mkqa":
        return f"MKQA · {LANGUAGE_NAMES[row['language']]}"
    return (f"{LEAGUE_NAMES.get(row['league'], row['league'])} · {LANGUAGE_NAMES[row['language']]}"
            f" · {CONDITION_NAMES.get(row['condition'], row['condition'])}")


def load_scored(path):
    path = Path(path)
    if not path.exists():
        sys.exit(f"Missing {path}. Run scripts/06_score.py first.")
    rows = read_csv(path)
    needed = set(GROUP_COLUMNS) | {"prompt_id", "question_id", "response", "score"}
    missing = needed - set(rows[0].keys()) if rows else needed
    if missing:
        sys.exit(f"{path.name} is missing the columns {sorted(missing)}. Is it a file from script 06?")
    return rows


def draw_sample(scored_rows):
    """20 replies per group (seed 440), shuffled. Failed requests (no score) are not sampled."""
    rng = random.Random(SEED)
    groups = {}
    for r in scored_rows:
        if r["score"] in ("0", "1"):
            groups.setdefault(tuple(r[c] for c in GROUP_COLUMNS), []).append(r)
    sample = []
    for key in sorted(groups):
        members = sorted(groups[key], key=lambda r: r["prompt_id"])
        if len(members) < PER_GROUP:
            print(f"  Note: group {key} has only {len(members)} scored replies; all are used.")
            sample.extend(members)
        else:
            sample.extend(rng.sample(members, PER_GROUP))
    rng.shuffle(sample)
    return sample


def row_height(values_and_widths):
    """Rough row height so wrapped text is visible when the file opens."""
    lines = 1
    for text, width in values_and_widths:
        chars_per_line = max(int(width * 0.9), 1)
        n = sum(max(1, -(-len(part) // chars_per_line)) for part in str(text).split("\n"))
        lines = max(lines, n)
    return min(15 * lines + 4, 160)


# ---------- make: build the blind sheet ----------

COLUMNS = [  # (header, width)
    ("#", 5),
    ("Group", 30),
    ("Question (as asked)", 44),
    ("Fact given in the prompt", 38),
    ("Answer key (English)", 18),
    ("Answer key (Malay)", 18),
    ("All accepted answers", 30),
    ("Chatbot reply", 44),
    (SCORE_HEADER, 17),
    ("Notes (e.g. hedged, spelling variant, key may be wrong)", 32),
    (ID_HEADER, 34),
]

RULES = [
    ("Correct", "The reply gives the answer key (or one of the accepted answers)."),
    ("Correct", "The right answer, but in a different language from the question (e.g. the reply is in "
                "English, or in a third language, when the question was in Malay). The reply language "
                "is recorded separately by the script."),
    ("Correct", "The right answer plus extra details, e.g. \"Marang, Terengganu, Malaysia\", as long as "
                "the fact that was asked is right."),
    ("Correct", "A different spelling that clearly names the same place, person or team, e.g. "
                "\"Kuala Trengganu\" for \"Kuala Terengganu\"."),
    ("Incorrect", "A nearby but different place, e.g. \"Kuala Terengganu\" when the answer is \"Marang\"."),
    ("Incorrect", "Only a broader area, e.g. \"Terengganu\" or \"Malaysia\" when the answer is \"Marang\"."),
    ("Incorrect", "Several possible answers (\"either X or Y\", \"X atau Y\"), even if one of them is "
                  "right: the reply does not commit to one answer."),
    ("Incorrect", "A refusal (\"I don't know\", \"tidak pasti\") or an empty reply."),
    ("Incorrect", "Any other wrong answer."),
]


def style_header(cell):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = BORDER


def build_instructions(ws, name, n_rows, last_row):
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 100
    rows = [
        (f"Blind hand check: {name}", None),
        ("DS 440: The Language Gap", None),
        (None, None),
        ("Purpose", "To measure how well the automatic scoring (scripts/06_score.py) works. Your scores are "
                    "compared with the script's scores using Cohen's kappa. The script's scores are NOT in "
                    "this file, so they cannot influence you."),
        ("What to do", f"Go to the \"{SHEET}\" sheet. For each of the {n_rows} replies, read the question, the "
                       "answer key and the chatbot reply, then choose correct or incorrect in the yellow "
                       f"\"{SCORE_HEADER}\" column. Use the yellow Notes column for anything unusual. "
                       "Replies are shown exactly as the chatbot sent them: ** marks bold text and can be ignored."),
        ("Done so far", f'=COUNTIF(\'{SHEET}\'!I2:I{last_row},"correct")+COUNTIF(\'{SHEET}\'!I2:I{last_row},"incorrect")'
                        f'&" of {n_rows} replies scored"'),
        (None, None),
        ("Rules", "From notes/scoring_rubric.md (\"Edge cases\", hand score column):"),
    ]
    rows += RULES
    rows += [
        ("Answer key", "Score each reply against the answer key and accepted answers shown, not against your "
                       "own memory. If you think the answer key itself is wrong, still score against it and "
                       "write \"key may be wrong\" in Notes."),
        (None, None),
        ("Example", "Not part of the sample; it only shows the format."),
        ("Question", "Where was the footballer Mohd Sharmiza Yusoff born?"),
        ("Answer key", "Marang"),
        ("Reply", "Kuala Terengganu, Malaysia."),
        ("Your score", "incorrect"),
        ("Notes", "nearby town, not Marang"),
        (None, None),
        ("Ground rules", "Do the scoring yourself, without AI help (no ChatGPT, Claude, Gemini in Colab or "
                         "Copilot in Excel). You may look things up, e.g. on Wikipedia or a map, to check a "
                         "spelling or a place. Do not open the scored file (data/scored/) or the results "
                         "while you are doing this."),
        ("Do not", f"Do not delete rows or edit the grey \"{ID_HEADER}\" column: it links each row back to "
                   "the automatic score. Sorting or filtering the rows is fine."),
        ("When done", "Save as .xlsx with the same file name, put it in data/hand_check/, then run: "
                      f"python scripts/07_hand_check.py kappa --scored data/scored/{name}_scored.csv"),
    ]
    for i, (label, text) in enumerate(rows, start=1):
        a, b = ws.cell(i, 1, label), ws.cell(i, 2, text)
        a.font = Font(name=FONT, bold=True)
        b.font = Font(name=FONT)
        a.alignment = Alignment(vertical="top")
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if label in ("Correct", "Incorrect"):
            a.font = Font(name=FONT, bold=True, color="375623" if label == "Correct" else "C00000")
        if label == "Your score" or label == "Notes":
            b.fill = INPUT_FILL
        if text is not None:
            ws.row_dimensions[i].height = row_height([(text, 100)]) if not str(text).startswith("=") else 18
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    ws["B6"].font = Font(name=FONT, bold=True)


def make(args):
    scored_path = Path(args.scored)
    name = name_of(scored_path)
    out = Path(args.out) if args.out else HAND_CHECK_DIR / f"{name}_hand_check.xlsx"
    if out.exists() and not args.overwrite:
        sys.exit(f"{out} already exists. It may already have your scores in it, so it was not replaced. "
                 "Use --overwrite only if you are sure.")
    if not QUESTIONS_PATH.exists():
        sys.exit(f"Missing {QUESTIONS_PATH}. Run scripts/04_combine_questions.py first.")
    questions = {q["question_id"]: q for q in read_csv(QUESTIONS_PATH)}

    sample = draw_sample(load_scored(scored_path))
    missing = [r["question_id"] for r in sample if r["question_id"] not in questions]
    if missing:
        sys.exit(f"{len(missing)} sampled replies have no question in questions_final.csv "
                 f"(e.g. {missing[0]}). Score the file again with scripts/06_score.py.")

    wb = Workbook()
    info = wb.active
    info.title = "How to use"
    ws = wb.create_sheet(SHEET)

    for j, (header, width) in enumerate(COLUMNS, start=1):
        style_header(ws.cell(1, j, header))
        ws.column_dimensions[ws.cell(1, j).column_letter].width = width
    ws.row_dimensions[1].height = 45

    for i, r in enumerate(sample, start=1):
        q = questions[r["question_id"]]
        lang = r["language"]
        fact = q[f"fact_{lang}"] if r["condition"] == "with_facts" else "(none)"
        accepted = " | ".join(a for a in q["accepted_answers"].split("|") if a)
        values = [i, group_label(r), q[f"question_{lang}"], fact, q["answer_en"], q["answer_ms"],
                  accepted, r["response"], None, None, r["prompt_id"]]
        for j, v in enumerate(values, start=1):
            c = ws.cell(i + 1, j, v)
            c.font = Font(name=FONT, color="595959" if j == len(values) else "000000")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
        for j in (9, 10):
            ws.cell(i + 1, j).fill = INPUT_FILL
        ws.cell(i + 1, len(values)).fill = ID_FILL
        ws.row_dimensions[i + 1].height = row_height(
            [(v, COLUMNS[j][1]) for j, v in enumerate(values) if v is not None and j not in (0, 10)])

    last = len(sample) + 1
    dv = DataValidation(type="list", formula1='"correct,incorrect"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Correct or incorrect",
                        error="Please choose correct or incorrect from the list.")
    ws.add_data_validation(dv)
    dv.add(f"I2:I{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:K{last}"

    build_instructions(info, name, len(sample), last)
    wb.active = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    counts = {}
    for r in sample:
        counts[group_label(r)] = counts.get(group_label(r), 0) + 1
    print(f"\n{scored_path.name}: {len(sample)} replies sampled (seed {SEED}), "
          f"{len(counts)} groups: " + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())))
    print(f"  Saved {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out} (no automatic scores in it)")


# ---------- kappa: compare the hand scores with the automatic scores ----------

def read_sheet(path):
    wb = load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"{path} has no sheet called \"{SHEET}\".")
    ws = wb[SHEET]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    for needed in ("#", SCORE_HEADER, ID_HEADER):
        if needed not in header:
            sys.exit(f"Column \"{needed}\" not found in the \"{SHEET}\" sheet. Was it renamed?")
    col = {h: header.index(h) for h in header if h}
    notes_col = next((i for h, i in col.items() if h.startswith("Notes")), None)
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        pid = values[col[ID_HEADER]]
        if not pid:
            continue
        raw = values[col[SCORE_HEADER]]
        text = str(raw).strip().lower() if raw is not None else ""
        hand = {"correct": 1, "incorrect": 0, "1": 1, "0": 0}.get(text)
        rows.append({"check": values[col["#"]], "prompt_id": str(pid).strip(), "hand_score": hand,
                     "hand_raw": text,
                     "notes": (values[notes_col] or "") if notes_col is not None else ""})
    return rows


def cohen_kappa(pairs):
    """pairs: list of (hand, auto) with values 0/1. Returns (n, agree, p_o, kappa or None)."""
    n = len(pairs)
    agree = sum(h == a for h, a in pairs)
    p_o = agree / n
    p_hand, p_auto = sum(h for h, _ in pairs) / n, sum(a for _, a in pairs) / n
    p_e = p_hand * p_auto + (1 - p_hand) * (1 - p_auto)
    kappa = None if p_e >= 1 else (p_o - p_e) / (1 - p_e)
    return n, agree, p_o, kappa


def kappa(args):
    scored_path = Path(args.scored)
    name = name_of(scored_path)
    sheet_path = Path(args.sheet) if args.sheet else HAND_CHECK_DIR / f"{name}_hand_check.xlsx"
    if not sheet_path.exists():
        sys.exit(f"Missing {sheet_path}. Make it with: python scripts/07_hand_check.py make --scored {args.scored}")

    scored_rows = load_scored(scored_path)
    scored = {r["prompt_id"]: r for r in scored_rows}
    hand_rows = read_sheet(sheet_path)

    unknown = [h["prompt_id"] for h in hand_rows if h["prompt_id"] not in scored]
    if unknown:
        sys.exit(f"{len(unknown)} prompt IDs in the sheet are not in {scored_path.name} (e.g. {unknown[0]}). "
                 "Is it the sheet for this chatbot?")
    bad = [h for h in hand_rows if h["hand_score"] is None]
    if bad:
        empty = [str(h["check"]) for h in bad if not h["hand_raw"]]
        other = [f"#{h['check']} ('{h['hand_raw']}')" for h in bad if h["hand_raw"]]
        msg = []
        if empty:
            msg.append(f"{len(empty)} {'reply has' if len(empty) == 1 else 'replies have'} no score yet: #"
                       + ", #".join(empty[:30])
                       + (" ..." if len(empty) > 30 else ""))
        if other:
            msg.append("these scores are not 'correct' or 'incorrect': " + ", ".join(other[:30]))
        sys.exit("Not finished: " + "; ".join(msg))

    expected = {r["prompt_id"] for r in draw_sample(scored_rows)}
    found = {h["prompt_id"] for h in hand_rows}
    if found != expected:
        print(f"  Warning: the sheet does not match the seed-{SEED} sample "
              f"({len(expected - found)} sampled replies missing, {len(found - expected)} extra). "
              "Kappa is computed for the rows in the sheet.")

    out_rows = []
    for h in sorted(hand_rows, key=lambda h: int(h["check"]) if str(h["check"]).isdigit() else 0):
        s = scored[h["prompt_id"]]
        auto = int(s["score"])
        out_rows.append({
            "check": h["check"], "prompt_id": h["prompt_id"], "question_id": s["question_id"],
            **{c: s[c] for c in GROUP_COLUMNS},
            "answer_en": s.get("answer_en", ""), "answer_ms": s.get("answer_ms", ""),
            "response": s["response"], "matched_answer": s.get("matched_answer", ""),
            "refusal": s.get("refusal", ""), "reply_language": s.get("reply_language", ""),
            "auto_score": auto, "hand_score": h["hand_score"], "agree": int(auto == h["hand_score"]),
            "notes": h["notes"],
        })

    summary = []

    def add(label, rows, key=None):
        n, agree, p_o, k = cohen_kappa([(r["hand_score"], r["auto_score"]) for r in rows])
        entry = {"group": label, **dict(zip(GROUP_COLUMNS, key or ("all",) * 4)), "n": n, "agree": agree,
                 "agreement": round(p_o, 3),
                 # kappa only for all 200: with 20 replies per group it is unstable, and it is 0 or not
                 # defined whenever one scorer marks every reply in the group correct
                 "kappa": "" if k is None or key else round(k, 3),
                 "auto_correct": sum(r["auto_score"] for r in rows),
                 "hand_correct": sum(r["hand_score"] for r in rows),
                 "auto_1_hand_0": sum(r["auto_score"] == 1 and r["hand_score"] == 0 for r in rows),
                 "auto_0_hand_1": sum(r["auto_score"] == 0 and r["hand_score"] == 1 for r in rows)}
        summary.append(entry)
        return entry

    overall = add("all", out_rows)
    by_group = {}
    for r in out_rows:
        by_group.setdefault(tuple(r[c] for c in GROUP_COLUMNS), []).append(r)
    for key in sorted(by_group):
        add(group_label(dict(zip(GROUP_COLUMNS, key))), by_group[key], key)

    detail_path = RESULTS_DIR / f"hand_check_{name}.csv"
    summary_path = RESULTS_DIR / f"hand_check_summary_{name}.csv"
    write_csv(detail_path, list(out_rows[0].keys()), out_rows)
    write_csv(summary_path, list(summary[0].keys()), summary)

    k = overall["kappa"]
    print(f"\n{sheet_path.name}: {overall['n']} replies hand-checked")
    print(f"  Agreement: {overall['agree']}/{overall['n']} = {overall['agreement']:.1%}")
    print(f"  Cohen's kappa: {k if k != '' else 'not defined (every reply got the same score from both)'}")
    print("\n  Confusion table          auto correct   auto incorrect")
    a1h1 = overall["auto_correct"] - overall["auto_1_hand_0"]
    a0h0 = overall["n"] - overall["auto_correct"] - overall["auto_0_hand_1"]
    print(f"    hand correct     {a1h1:>12}   {overall['auto_0_hand_1']:>14}")
    print(f"    hand incorrect   {overall['auto_1_hand_0']:>12}   {a0h0:>14}")
    print("\n  Agreement per group:")
    for s in summary[1:]:
        print(f"    {s['group']:46} {s['agree']:>2}/{s['n']:<2}  (auto correct {s['auto_correct']:>2}, "
              f"hand correct {s['hand_correct']:>2})")
    disagreements = [r for r in out_rows if not r["agree"]]
    print(f"\n  Disagreements: {len(disagreements)}")
    for r in disagreements:
        print(f"    #{r['check']} {r['prompt_id']}: auto {r['auto_score']}, hand {r['hand_score']} | "
              f"answer {r['answer_en']} | reply: {r['response'][:70]!r}"
              + (f" | notes: {r['notes']}" if r["notes"] else ""))
    print(f"\n  Saved {detail_path.relative_to(ROOT)} and {summary_path.relative_to(ROOT)}")


def main():
    parser = argparse.ArgumentParser(description="Blind hand check of the automatic scores (Cohen's kappa)")
    sub = parser.add_subparsers(dest="mode", required=True)
    m = sub.add_parser("make", help="make the blind hand-check sheet")
    m.add_argument("--scored", required=True, help="scored file from script 06")
    m.add_argument("--out", help="where to save the sheet (default: data/hand_check/<name>_hand_check.xlsx)")
    m.add_argument("--overwrite", action="store_true", help="replace an existing sheet (careful: loses your scores)")
    k = sub.add_parser("kappa", help="compare the filled sheet with the automatic scores")
    k.add_argument("--scored", required=True, help="scored file from script 06")
    k.add_argument("--sheet", help="filled sheet (default: data/hand_check/<name>_hand_check.xlsx)")
    args = parser.parse_args()
    make(args) if args.mode == "make" else kappa(args)


if __name__ == "__main__":
    main()
